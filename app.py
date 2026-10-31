from __future__ import annotations

import csv
import html
import hashlib
import hmac
import io
import json
import os
import re
import sys
import threading
import webbrowser
from collections import defaultdict
from datetime import datetime
from email.parser import BytesParser
from email.policy import default
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import Iterable
from urllib import error, parse, request

import xlrd
from openpyxl import Workbook, load_workbook


HOST = "127.0.0.1"
PORT = 8000
RUNTIME_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
BUNDLE_DIR = getattr(sys, "_MEIPASS", RUNTIME_DIR)
ASSET_TEMPLATE = os.path.join(BUNDLE_DIR, "assets", "inventory_export_1.csv")
BUTLER_ASSET_TEMPLATE = os.path.join(BUNDLE_DIR, "assets", "merged_inventory_export.csv")
SHIPPING_DATA_PATH = os.path.join(BUNDLE_DIR, "assets", "shipping_info.csv")
ORDER_NOTES_PATH = os.path.join(RUNTIME_DIR, "order_notes.json")
BUTLER_EXTERNAL_TEMPLATE = os.path.join(
    os.path.dirname(os.path.dirname(BUNDLE_DIR)),
    "test FOA",
    "Butler",
    "merged_inventory_export.csv",
)
TARGET_LOCATION = "Furniture of America"
NOT_STOCKED = "not stocked"
MIN_EXPORT_QTY = 5
ENV_PATH = os.path.join(RUNTIME_DIR, ".env")
SIBLING_ENV_PATH = os.path.join(
    os.path.dirname(RUNTIME_DIR),
    "Gooddegg_order_management",
    ".env",
)
SHOPIFY_DEFAULT_API_VERSION = "2025-10"
SHOPIFY_SCOPES = "read_orders"
SHOPIFY_STATE_KEY = "shopify_oauth_state"
SHOPIFY_TOKEN_KEY = "SHOPIFY_ADMIN_ACCESS_TOKEN"
SHOPIFY_ORDERS_LIMIT = 20
REMOVED_OUTPUT_COLUMNS = {
    "Incoming (not editable)",
    "Unavailable (not editable)",
    "Committed (not editable)",
    "Available (not editable)",
    "On hand (current)",
}

SKU_KEYWORDS = ("sku", "itemcode", "item code", "item", "product code", "product")
QTY_EXACT = ("qty", "quantity", "qty now", "stock qty", "available qty", "on hand qty")
QTY_PARTIAL = ("qty", "quantity", "stock", "on hand", "available")
SHOPIFY_RUNTIME_STATE: dict[str, str] = {}
SHIPPING_RECORDS_CACHE: dict[str, dict[str, str]] | None = None
ORDER_NOTES_LOCK = threading.Lock()


def load_shipping_records() -> dict[str, dict[str, str]]:
    global SHIPPING_RECORDS_CACHE
    if SHIPPING_RECORDS_CACHE is not None:
        return SHIPPING_RECORDS_CACHE

    if not os.path.exists(SHIPPING_DATA_PATH):
        raise ValueError("Shipping information data file was not found.")

    records: dict[str, dict[str, str]] = {}
    with open(SHIPPING_DATA_PATH, encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            sku = (row.get("sku") or "").strip().upper()
            if sku:
                records[sku] = {key: (value or "").strip() for key, value in row.items()}
    SHIPPING_RECORDS_CACHE = records
    return records


def find_shipping_records(query: str) -> tuple[list[dict[str, str]], list[str]]:
    normalized_query = query.strip().upper()
    if not normalized_query:
        return [], []

    records = load_shipping_records()
    exact = records.get(normalized_query)
    if exact is None:
        suggestions = [sku for sku in records if normalized_query in sku][:10]
        return [], suggestions

    results = [exact]
    part_skus = [part.strip().upper() for part in exact.get("parts", "").split(",") if part.strip()]
    results.extend(records[sku] for sku in part_skus if sku in records)
    return results, []


def _read_order_notes() -> list[dict[str, str]]:
    if not os.path.exists(ORDER_NOTES_PATH):
        return []
    with open(ORDER_NOTES_PATH, encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, list):
        raise ValueError("Order notes data file is invalid.")
    return [note for note in data if isinstance(note, dict)]


def load_order_notes() -> list[dict[str, str]]:
    with ORDER_NOTES_LOCK:
        return _read_order_notes()


def _write_order_notes(notes: list[dict[str, str]]) -> None:
    temporary_path = f"{ORDER_NOTES_PATH}.tmp"
    with open(temporary_path, "w", encoding="utf-8") as handle:
        json.dump(notes, handle, ensure_ascii=False, indent=2)
    os.replace(temporary_path, ORDER_NOTES_PATH)


def save_order_note(note_id: str, order_number: str, note_text: str) -> None:
    clean_order_number = order_number.strip()[:200]
    clean_note_text = note_text.strip()[:10000]
    if not clean_note_text:
        raise ValueError("Please write a note before saving.")

    now = datetime.now().strftime("%d %b %Y, %I:%M %p")
    with ORDER_NOTES_LOCK:
        notes = _read_order_notes()
        existing = next((note for note in notes if note.get("id") == note_id), None) if note_id else None
        if existing is not None:
            existing["order_number"] = clean_order_number
            existing["note"] = clean_note_text
            existing["updated_at"] = now
        else:
            notes.insert(
                0,
                {
                    "id": hashlib.sha256(os.urandom(32)).hexdigest()[:16],
                    "order_number": clean_order_number,
                    "note": clean_note_text,
                    "created_at": now,
                    "updated_at": now,
                },
            )
        _write_order_notes(notes)


def delete_order_note(note_id: str) -> None:
    with ORDER_NOTES_LOCK:
        notes = _read_order_notes()
        filtered_notes = [note for note in notes if note.get("id") != note_id]
        if len(filtered_notes) == len(notes):
            raise ValueError("Order note was not found.")
        _write_order_notes(filtered_notes)


def load_dotenv_file(path: str) -> dict[str, str]:
    if not os.path.exists(path):
        return {}

    values: dict[str, str] = {}
    with open(path, encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue

            key, value = line.split("=", 1)
            values[key.strip()] = value.strip()

    return values


def load_shopify_settings() -> dict[str, str]:
    settings: dict[str, str] = {}
    for candidate in (SIBLING_ENV_PATH, ENV_PATH):
        settings.update(load_dotenv_file(candidate))

    return {
        "shop_domain": settings.get("SHOPIFY_SHOP_DOMAIN", "").strip(),
        "api_version": settings.get("SHOPIFY_API_VERSION", SHOPIFY_DEFAULT_API_VERSION).strip() or SHOPIFY_DEFAULT_API_VERSION,
        "client_id": settings.get("SHOPIFY_CLIENT_ID", "").strip(),
        "client_secret": settings.get("SHOPIFY_CLIENT_SECRET", "").strip(),
        "admin_access_token": settings.get(SHOPIFY_TOKEN_KEY, "").strip(),
    }


def save_env_value(path: str, key: str, value: str) -> None:
    lines: list[str] = []
    if os.path.exists(path):
        with open(path, encoding="utf-8") as handle:
            lines = handle.read().splitlines()

    updated = False
    for index, line in enumerate(lines):
        if line.startswith(f"{key}="):
            lines[index] = f"{key}={value}"
            updated = True
            break

    if not updated:
        lines.append(f"{key}={value}")

    content = "\n".join(lines).rstrip() + "\n"
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(content)


def normalize_shop_domain(value: str) -> str:
    sanitized = value.strip()
    sanitized = sanitized.replace("https://", "")
    sanitized = sanitized.replace("http://", "")
    sanitized = sanitized.strip("/")
    if sanitized and "." not in sanitized:
        sanitized = f"{sanitized}.myshopify.com"
    return sanitized


def shopify_redirect_uri() -> str:
    return f"http://{HOST}:{PORT}/shopify/callback"


def build_shopify_authorization_url(settings: dict[str, str], state: str) -> str:
    params = parse.urlencode(
        {
            "client_id": settings["client_id"],
            "scope": SHOPIFY_SCOPES,
            "redirect_uri": shopify_redirect_uri(),
            "state": state,
        }
    )
    return f"https://{normalize_shop_domain(settings['shop_domain'])}/admin/oauth/authorize?{params}"


def validate_shopify_hmac(query_params: dict[str, str], client_secret: str) -> bool:
    received_hmac = query_params.get("hmac", "").strip().lower()
    if not received_hmac:
        return False

    message_parts = []
    for key in sorted(query_params):
        if key in {"hmac", "signature"}:
            continue
        message_parts.append(f"{key}={query_params[key]}")

    message = "&".join(message_parts).encode("utf-8")
    digest = hmac.new(client_secret.encode("utf-8"), message, hashlib.sha256).hexdigest().lower()
    return hmac.compare_digest(digest, received_hmac)


def exchange_shopify_code_for_token(settings: dict[str, str], code: str) -> str:
    url = f"https://{normalize_shop_domain(settings['shop_domain'])}/admin/oauth/access_token"
    payload = parse.urlencode(
        {
            "client_id": settings["client_id"],
            "client_secret": settings["client_secret"],
            "code": code,
        }
    ).encode("utf-8")

    req = request.Request(url, data=payload, method="POST")
    req.add_header("Content-Type", "application/x-www-form-urlencoded")
    try:
        with request.urlopen(req, timeout=30) as response:
            body = response.read().decode("utf-8")
    except error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise ValueError(f"Token exchange failed: {details or exc.reason}") from exc

    data = json.loads(body)
    token = str(data.get("access_token", "")).strip()
    if not token:
        raise ValueError("Shopify response did not include an access token.")
    return token


def fetch_shopify_orders(settings: dict[str, str], limit: int = SHOPIFY_ORDERS_LIMIT) -> list[dict[str, str]]:
    token = settings.get("admin_access_token", "").strip()
    if not token:
        raise ValueError("Shopify access token is missing.")

    domain = normalize_shop_domain(settings["shop_domain"])
    api_version = settings.get("api_version", SHOPIFY_DEFAULT_API_VERSION) or SHOPIFY_DEFAULT_API_VERSION
    url = f"https://{domain}/admin/api/{api_version}/orders.json?status=any&limit={max(1, min(limit, 250))}"

    req = request.Request(url, method="GET")
    req.add_header("Accept", "application/json")
    req.add_header("X-Shopify-Access-Token", token)

    try:
        with request.urlopen(req, timeout=30) as response:
            body = response.read().decode("utf-8")
    except error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise ValueError(f"Shopify API error: {details or exc.reason}") from exc

    payload = json.loads(body)
    orders = payload.get("orders", [])
    result: list[dict[str, str]] = []
    for order in orders:
        customer = order.get("customer") or {}
        first_name = str(customer.get("first_name", "")).strip()
        last_name = str(customer.get("last_name", "")).strip()
        customer_name = f"{first_name} {last_name}".strip() or "N/A"
        result.append(
            {
                "name": str(order.get("name", "")).strip() or "N/A",
                "customer": customer_name,
                "total_price": str(order.get("total_price", "")).strip() or "N/A",
                "currency": str(order.get("currency", "")).strip() or "",
                "financial_status": str(order.get("financial_status", "")).strip() or "N/A",
                "fulfillment_status": str(order.get("fulfillment_status", "")).strip() or "unfulfilled",
                "created_at": str(order.get("created_at", "")).strip() or "N/A",
            }
        )
    return result


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def clean_sku(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text


def normalize_numeric_sku(value: object) -> str:
    sku = clean_sku(value)
    if sku.isdigit():
        return sku.lstrip("0") or "0"
    return sku


def parse_quantity(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0

    try:
        return float(text)
    except ValueError:
        return 0.0


def score_sku_header(header: str) -> int:
    if not header:
        return 0
    if header in SKU_KEYWORDS:
        return 100
    if any(keyword in header for keyword in SKU_KEYWORDS):
        return 50
    return 0


def score_qty_header(header: str) -> int:
    if not header:
        return 0
    if header in QTY_EXACT:
        return 100
    if "qty now" in header:
        return 95
    if any(keyword in header for keyword in QTY_PARTIAL):
        score = 60
        if "w" in header and "qty" in header:
            score -= 20
        if "coming" in header or "week" in header or "eta" in header:
            score -= 25
        return score
    return 0


def build_combined_headers(rows: list[list[object]], header_end_index: int) -> list[str]:
    max_cols = max((len(row) for row in rows[: header_end_index + 1]), default=0)
    combined_headers: list[str] = []

    for col_index in range(max_cols):
        pieces: list[str] = []
        for row_index in range(header_end_index + 1):
            row = rows[row_index]
            value = normalize_text(row[col_index] if col_index < len(row) else "")
            if value and value not in pieces:
                pieces.append(value)
        combined_headers.append(" ".join(pieces).strip())

    return combined_headers


def header_score(row: list[object]) -> int:
    row_score = 0
    for cell in row:
        text = normalize_text(cell)
        row_score += score_sku_header(text)
        row_score += score_qty_header(text)
    return row_score


def detect_header_row(rows: list[list[object]]) -> int:
    best_index = 0
    best_score = -1

    limit = min(len(rows), 12)
    for row_index in range(limit):
        row_score = header_score(rows[row_index])
        if row_score > best_score:
            best_score = row_score
            best_index = row_index

    return best_index


def detect_header_block_end(rows: list[list[object]]) -> int:
    limit = min(len(rows), 12)
    scored_rows = [index for index in range(limit) if header_score(rows[index]) > 0]
    if scored_rows:
        return scored_rows[-1]
    return detect_header_row(rows)


def detect_columns(headers: list[str]) -> tuple[int, int]:
    sku_index = -1
    sku_score = -1
    qty_index = -1
    qty_score = -1

    for index, header in enumerate(headers):
        current_sku_score = score_sku_header(header)
        current_qty_score = score_qty_header(header)

        if current_sku_score > sku_score:
            sku_score = current_sku_score
            sku_index = index

        if current_qty_score > qty_score:
            qty_score = current_qty_score
            qty_index = index

    if sku_score <= 0:
        raise ValueError("Could not find an SKU column.")
    if qty_score <= 0:
        raise ValueError("Could not find a quantity column.")

    return sku_index, qty_index


def aggregate_rows(rows: Iterable[list[object]], sku_index: int, qty_index: int) -> dict[str, float]:
    totals: dict[str, float] = defaultdict(float)

    for row in rows:
        if sku_index >= len(row):
            continue

        sku = clean_sku(row[sku_index])
        if not sku:
            continue

        qty_value = row[qty_index] if qty_index < len(row) else None
        qty = parse_quantity(qty_value)
        totals[sku] += qty

    return totals


def read_csv_inventory(file_bytes: bytes) -> dict[str, float]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            decoded = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("CSV file encoding is not supported.")

    all_rows = list(csv.reader(io.StringIO(decoded)))
    if not all_rows:
        return {}

    header_end_index = detect_header_block_end(all_rows)
    headers = build_combined_headers(all_rows[: header_end_index + 1], header_end_index)
    sku_index, qty_index = detect_columns(headers)
    data_rows = all_rows[header_end_index + 1 :]
    return aggregate_rows(data_rows, sku_index, qty_index)


def read_xlsx_inventory(file_bytes: bytes) -> dict[str, float]:
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)

    for worksheet in workbook.worksheets:
        preview_rows: list[list[object]] = []
        all_rows: list[list[object]] = []

        for row in worksheet.iter_rows(values_only=True):
            row_values = list(row)
            all_rows.append(row_values)
            if len(preview_rows) < 12:
                preview_rows.append(row_values)

        if not any(any(cell is not None and str(cell).strip() for cell in row) for row in all_rows):
            continue

        header_end_index = detect_header_block_end(preview_rows)
        headers = build_combined_headers(preview_rows, header_end_index)
        sku_index, qty_index = detect_columns(headers)
        data_rows = all_rows[header_end_index + 1 :]
        return aggregate_rows(data_rows, sku_index, qty_index)

    return {}


def read_xls_inventory(file_bytes: bytes) -> dict[str, float]:
    workbook = xlrd.open_workbook(file_contents=file_bytes)

    for sheet in workbook.sheets():
        all_rows = [sheet.row_values(row_index) for row_index in range(sheet.nrows)]
        if not all_rows:
            continue

        preview_rows = all_rows[:12]
        header_end_index = detect_header_block_end(preview_rows)
        headers = build_combined_headers(preview_rows, header_end_index)
        sku_index, qty_index = detect_columns(headers)
        data_rows = all_rows[header_end_index + 1 :]
        return aggregate_rows(data_rows, sku_index, qty_index)

    return {}


def read_inventory_file(filename: str, file_bytes: bytes) -> dict[str, float]:
    extension = os.path.splitext(filename.lower())[1]

    if extension == ".csv":
        return read_csv_inventory(file_bytes)
    if extension == ".xlsx":
        return read_xlsx_inventory(file_bytes)
    if extension == ".xls":
        return read_xls_inventory(file_bytes)

    raise ValueError(f"Unsupported file type: {extension or filename}")


def merge_inventory(files: list[tuple[str, bytes]]) -> dict[str, float]:
    grand_totals: dict[str, float] = defaultdict(float)

    for filename, file_bytes in files:
        file_totals = read_inventory_file(filename, file_bytes)
        for sku, qty in file_totals.items():
            grand_totals[sku] += qty

    return dict(sorted(grand_totals.items()))


def format_qty_for_export(qty: float) -> str:
    if float(qty).is_integer():
        return str(int(qty))
    return str(qty)


def decode_csv_bytes(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV file encoding is not supported.")


def load_asset_rows(asset_file: tuple[str, bytes] | None) -> tuple[list[dict[str, str]], list[str]]:
    if asset_file is None:
        if not os.path.exists(ASSET_TEMPLATE):
            raise ValueError(f"Asset template not found: {ASSET_TEMPLATE}")
        with open(ASSET_TEMPLATE, newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            rows = list(reader)
            fieldnames = reader.fieldnames
    else:
        filename, file_bytes = asset_file
        if os.path.splitext(filename.lower())[1] != ".csv":
            raise ValueError("The asset export file must be a CSV.")
        reader = csv.DictReader(io.StringIO(decode_csv_bytes(file_bytes)))
        rows = list(reader)
        fieldnames = reader.fieldnames

    if not fieldnames:
        raise ValueError("Asset template has no header row.")

    output_fieldnames = [name for name in fieldnames if name not in REMOVED_OUTPUT_COLUMNS]
    return rows, output_fieldnames


def build_updated_inventory_csv(
    totals: dict[str, float],
    asset_file: tuple[str, bytes] | None = None,
    template_path: str | None = None,
    target_location: str | None = TARGET_LOCATION,
    minimum_quantity: float = MIN_EXPORT_QTY,
    normalize_numeric_skus: bool = False,
) -> bytes:
    if asset_file is None and template_path is not None:
        with open(template_path, "rb") as handle:
            asset_file = (os.path.basename(template_path), handle.read())
    rows, fieldnames = load_asset_rows(asset_file)

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()

    quantity_totals: dict[str, float] = totals
    if normalize_numeric_skus:
        normalized_totals: dict[str, float] = defaultdict(float)
        for sku, quantity in totals.items():
            normalized_totals[normalize_numeric_sku(sku)] += quantity
        quantity_totals = normalized_totals

    for row in rows:
        location = clean_sku(row.get("Location"))
        if target_location is not None and location != target_location:
            continue

        sku = (
            normalize_numeric_sku(row.get("SKU"))
            if normalize_numeric_skus
            else clean_sku(row.get("SKU"))
        )
        available_value = normalize_text(row.get("Available (not editable)"))
        if target_location is None or available_value != NOT_STOCKED:
            quantity = quantity_totals.get(sku, 0.0)
            if quantity < minimum_quantity:
                quantity = 0.0
            row["On hand (new)"] = format_qty_for_export(quantity)
        writer.writerow({key: row.get(key, "") for key in fieldnames})

    return output.getvalue().encode("utf-8-sig")


def build_summary_workbook(totals: dict[str, float]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SKU Totals"
    worksheet.append(["SKU", "Total Qty"])

    for sku, qty in totals.items():
        if float(qty).is_integer():
            worksheet.append([sku, int(qty)])
        else:
            worksheet.append([sku, qty])

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 14

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.read()


def render_home(message: str = "", mode: str = "foa") -> bytes:
    escaped_message = html.escape(message)
    is_butler = mode == "butler"
    page_title = "Butler Inventory" if is_butler else "FOA Inventory"
    intro = (
        "Upload the Butler specialty product inventory CSV. Product is matched to SKU and Available is written to On hand (new)."
        if is_butler
        else "Upload your report files and generate the final CSV for Furniture of America. Mixed <strong>CSV</strong>, <strong>XLSX</strong>, and <strong>XLS</strong> files are supported."
    )
    upload_help = "Select the Butler specialty product inventory CSV." if is_butler else "Select two or more inventory reports."
    upload_name = "butler_inventory_file" if is_butler else "inventory_files"
    upload_multiple = "" if is_butler else " multiple"
    action = "/butler" if is_butler else "/"
    template_label = "assets/merged_inventory_export.csv" if is_butler else "assets/inventory_export_1.csv"
    html_body = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Inventory Process Hub</title>
  <style>
    :root {{
      --bg: #f6f7f4;
      --panel: rgba(255, 255, 255, 0.82);
      --text: #17201b;
      --muted: #66736c;
      --line: rgba(23, 32, 27, 0.08);
      --accent: #163f31;
      --accent-soft: #edf3ef;
      --error-bg: #fff3f2;
      --error-text: #9a3b35;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: "Segoe UI", Tahoma, sans-serif;
      background:
        radial-gradient(circle at top left, rgba(22, 63, 49, 0.08), transparent 24%),
        radial-gradient(circle at bottom right, rgba(22, 63, 49, 0.06), transparent 20%),
        linear-gradient(180deg, #f8faf8 0%, #f3f5f2 100%);
      color: var(--text);
    }}
    .wrap {{
      min-height: 100vh;
      display: grid;
      grid-template-columns: 230px minmax(0, 720px);
      align-items: center;
      justify-content: center;
      gap: 24px;
      padding: 32px 18px;
    }}
    .sidebar {{
      align-self: stretch;
      min-height: 620px;
      padding: 24px 14px;
      border: 1px solid rgba(255,255,255,0.7);
      border-radius: 24px;
      background: rgba(255,255,255,0.7);
      box-shadow: 0 18px 50px rgba(28, 42, 35, 0.06);
      backdrop-filter: blur(14px);
    }}
    .brand {{
      padding: 8px 12px 24px;
      color: var(--accent);
      font-size: 1.05rem;
      font-weight: 700;
      letter-spacing: -0.02em;
    }}
    .menu-title {{
      padding: 0 12px 10px;
      color: var(--muted);
      font-size: 0.72rem;
      font-weight: 700;
      letter-spacing: 0.1em;
      text-transform: uppercase;
    }}
    .menu-item {{
      display: flex;
      align-items: center;
      gap: 10px;
      padding: 12px;
      border-radius: 14px;
      color: var(--accent);
      font-weight: 600;
      text-decoration: none;
    }}
    .menu-item.active {{
      background: var(--accent-soft);
    }}
    .menu-icon {{
      display: grid;
      width: 28px;
      height: 28px;
      place-items: center;
      border-radius: 9px;
      background: var(--accent);
      color: white;
      font-size: 0.78rem;
    }}
    .content {{
      width: 100%;
    }}
    .panel {{
      width: min(720px, 100%);
      background: var(--panel);
      border: 1px solid rgba(255,255,255,0.65);
      border-radius: 28px;
      padding: 34px;
      backdrop-filter: blur(14px);
      box-shadow: 0 18px 50px rgba(28, 42, 35, 0.08);
    }}
    .label {{
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 7px 12px;
      border-radius: 999px;
      background: var(--accent-soft);
      color: var(--accent);
      font-size: 0.78rem;
      font-weight: 600;
      letter-spacing: 0.08em;
      text-transform: uppercase;
    }}
    h1 {{
      margin: 18px 0 10px;
      font-size: clamp(2rem, 5vw, 3rem);
      line-height: 1.02;
      letter-spacing: -0.04em;
    }}
    .intro {{
      margin: 0 0 24px;
      color: var(--muted);
      font-size: 1rem;
      line-height: 1.65;
    }}
    .note {{
      margin-bottom: 18px;
      padding: 14px 16px;
      border-radius: 16px;
      border: 1px solid var(--line);
      background: rgba(255, 255, 255, 0.7);
      color: var(--muted);
      line-height: 1.55;
    }}
    .error {{
      background: var(--error-bg);
      color: var(--error-text);
    }}
    form {{
      display: grid;
      gap: 16px;
    }}
    .upload-zone {{
      border: 1px dashed rgba(23, 32, 27, 0.12);
      background: rgba(255, 255, 255, 0.72);
      border-radius: 22px;
      padding: 20px;
    }}
    .upload-meta {{
      display: grid;
      gap: 4px;
      margin-bottom: 12px;
    }}
    .upload-meta strong {{
      font-size: 1rem;
    }}
    .upload-meta span {{
      color: var(--muted);
      font-size: 0.94rem;
    }}
    input[type=file] {{
      width: 100%;
      border: 1px solid var(--line);
      border-radius: 14px;
      padding: 16px;
      background: white;
      font-size: 0.96rem;
    }}
    button {{
      border: 0;
      border-radius: 16px;
      padding: 15px 20px;
      background: var(--accent);
      color: white;
      font-size: 0.98rem;
      font-weight: 600;
      cursor: pointer;
      box-shadow: 0 10px 24px rgba(22, 63, 49, 0.16);
      transition: transform 160ms ease, box-shadow 160ms ease;
    }}
    button:hover {{
      transform: translateY(-1px);
      box-shadow: 0 14px 28px rgba(22, 63, 49, 0.2);
    }}
    .meta {{
      margin-top: 18px;
      padding-top: 16px;
      border-top: 1px solid var(--line);
      color: var(--muted);
      font-size: 0.95rem;
      line-height: 1.55;
    }}
    @media (max-width: 720px) {{
      .wrap {{
        grid-template-columns: 1fr;
        align-items: start;
        padding: 18px 14px;
      }}
      .sidebar {{
        min-height: auto;
      }}
      .panel {{
        padding: 24px;
        border-radius: 22px;
      }}
    }}
  </style>
</head>
<body>
  <div class="wrap">
    <aside class="sidebar">
      <div class="brand">Gooddegg Management</div>
      <div class="menu-title">Menu</div>
      <a class="menu-item {'active' if not is_butler else ''}" href="/">
        <span class="menu-icon">F</span>
        <span>FOA Inventory</span>
      </a>
      <a class="menu-item {'active' if is_butler else ''}" href="/butler">
        <span class="menu-icon">B</span>
        <span>Butler Inventory</span>
      </a>
      <a class="menu-item" href="/shipping">
        <span class="menu-icon">S</span>
        <span>Shipping Info</span>
      </a>
      <a class="menu-item" href="/order-notes">
        <span class="menu-icon">N</span>
        <span>Order Notes</span>
      </a>
    </aside>
    <main class="content">
      <div class="panel">
        <span class="label">Inventory Process</span>
        <h1>{page_title}</h1>
        <p class="intro">{intro}</p>
        <div class="note {('error' if message else '')}">{escaped_message or 'Stocked SKUs are updated from merged totals. Rows marked not stocked stay blank.'}</div>
        <form method="post" action="{action}" enctype="multipart/form-data">
          <div class="upload-zone">
            <div class="upload-meta">
              <strong>Source files</strong>
              <span>{upload_help}</span>
            </div>
            <input type="file" name="{upload_name}" accept=".csv,.xlsx,.xls"{upload_multiple} required>
          </div>
          <button type="submit">Generate {'Butler' if is_butler else 'Updated'} Export</button>
        </form>
        <div class="meta">Default base file: <strong>{template_label}</strong><br><a href="/orders">Open Shopify order list</a></div>
      </div>
    </main>
  </div>
</body>
</html>"""
    return html_body.encode("utf-8")


def render_shipping_page(query: str = "", error_message: str = "") -> bytes:
    records: list[dict[str, str]] = []
    suggestions: list[str] = []
    if query and not error_message:
        try:
            records, suggestions = find_shipping_records(query)
        except Exception as exc:
            error_message = str(exc)

    result_html = ""
    if records:
        cards = []
        for index, record in enumerate(records):
            sku = html.escape(record.get("sku", ""))
            description = html.escape(record.get("description", "") or record.get("name", ""))
            badge = "Selected SKU" if index == 0 else "Package / Component"

            def display(field: str, suffix: str = "") -> str:
                value = record.get(field, "")
                return f"{html.escape(value)}{suffix}" if value else '<span class="missing">Not provided</span>'

            cards.append(
                f"""<article class="result-card">
                  <div class="result-head">
                    <div><span class="result-badge">{badge}</span><h2>{sku}</h2></div>
                    <span class="item-type">{html.escape(record.get('item_type', ''))}</span>
                  </div>
                  <p class="description">{description}</p>
                  <div class="measurements">
                    <div><span>Length</span><strong>{display('length', ' in')}</strong></div>
                    <div><span>Width</span><strong>{display('width', ' in')}</strong></div>
                    <div><span>Height</span><strong>{display('height', ' in')}</strong></div>
                    <div><span>Weight</span><strong>{display('weight', ' lb')}</strong></div>
                  </div>
                </article>"""
            )
        result_html = '<section class="results">' + "".join(cards) + "</section>"
    elif query and suggestions:
        links = "".join(
            f'<a href="/shipping?sku={parse.quote(sku)}">{html.escape(sku)}</a>' for sku in suggestions
        )
        result_html = (
            '<div class="empty"><strong>No exact SKU found.</strong>'
            f'<p>Did you mean one of these?</p><div class="suggestions">{links}</div></div>'
        )
    elif query:
        result_html = '<div class="empty"><strong>No matching SKU found.</strong><p>Check the SKU and try again.</p></div>'

    error_html = f'<div class="error">{html.escape(error_message)}</div>' if error_message else ""
    html_body = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Shipping Info | Gooddegg Management</title>
  <style>
    :root {{ --bg:#f6f7f4; --panel:rgba(255,255,255,.86); --text:#17201b; --muted:#66736c; --line:rgba(23,32,27,.09); --accent:#163f31; --soft:#edf3ef; }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; font-family:"Segoe UI",Tahoma,sans-serif; color:var(--text); background:radial-gradient(circle at top left,rgba(22,63,49,.08),transparent 24%),linear-gradient(180deg,#f8faf8,#f3f5f2); }}
    .wrap {{ min-height:100vh; display:grid; grid-template-columns:230px minmax(0,850px); justify-content:center; gap:24px; padding:32px 18px; }}
    .sidebar {{ padding:24px 14px; border:1px solid rgba(255,255,255,.7); border-radius:24px; background:rgba(255,255,255,.7); box-shadow:0 18px 50px rgba(28,42,35,.06); }}
    .brand {{ padding:8px 12px 24px; color:var(--accent); font-size:1.05rem; font-weight:700; }}
    .menu-title {{ padding:0 12px 10px; color:var(--muted); font-size:.72rem; font-weight:700; letter-spacing:.1em; text-transform:uppercase; }}
    .menu-item {{ display:flex; align-items:center; gap:10px; padding:12px; border-radius:14px; color:var(--accent); font-weight:600; text-decoration:none; }}
    .menu-item.active {{ background:var(--soft); }}
    .menu-icon {{ display:grid; width:28px; height:28px; place-items:center; border-radius:9px; background:var(--accent); color:white; font-size:.78rem; }}
    main {{ min-width:0; }}
    .panel {{ background:var(--panel); border:1px solid rgba(255,255,255,.65); border-radius:28px; padding:34px; box-shadow:0 18px 50px rgba(28,42,35,.08); }}
    .label,.result-badge {{ display:inline-flex; padding:7px 12px; border-radius:999px; background:var(--soft); color:var(--accent); font-size:.76rem; font-weight:700; letter-spacing:.06em; text-transform:uppercase; }}
    h1 {{ margin:18px 0 10px; font-size:clamp(2rem,5vw,3rem); line-height:1.02; letter-spacing:-.04em; }}
    .intro,.description {{ color:var(--muted); line-height:1.6; }}
    form {{ display:flex; gap:10px; margin-top:24px; }}
    input {{ flex:1; min-width:0; border:1px solid var(--line); border-radius:16px; padding:15px 17px; background:white; color:var(--text); font-size:1rem; text-transform:uppercase; outline:none; }}
    input:focus {{ border-color:rgba(22,63,49,.45); box-shadow:0 0 0 4px rgba(22,63,49,.08); }}
    button {{ border:0; border-radius:16px; padding:15px 22px; background:var(--accent); color:white; font-size:.98rem; font-weight:700; cursor:pointer; }}
    .results {{ display:grid; gap:14px; margin-top:20px; }}
    .result-card,.empty,.error {{ border:1px solid var(--line); border-radius:20px; padding:20px; background:rgba(255,255,255,.76); }}
    .result-head {{ display:flex; align-items:start; justify-content:space-between; gap:14px; }}
    .result-head h2 {{ margin:10px 0 0; font-size:1.35rem; }}
    .item-type {{ color:var(--muted); font-size:.85rem; }}
    .description {{ margin:8px 0 16px; }}
    .measurements {{ display:grid; grid-template-columns:repeat(4,1fr); gap:10px; }}
    .measurements div {{ padding:14px; border-radius:14px; background:var(--soft); }}
    .measurements span {{ display:block; margin-bottom:5px; color:var(--muted); font-size:.78rem; }}
    .measurements strong {{ font-size:1.05rem; }}
    .measurements .missing {{ margin:0; font-size:.83rem; font-weight:500; }}
    .empty,.error {{ margin-top:20px; color:var(--muted); }}
    .error {{ background:#fff3f2; color:#9a3b35; }}
    .suggestions {{ display:flex; flex-wrap:wrap; gap:8px; }}
    .suggestions a {{ padding:7px 10px; border-radius:10px; background:var(--soft); color:var(--accent); text-decoration:none; font-weight:600; }}
    @media(max-width:760px) {{ .wrap {{ grid-template-columns:1fr; padding:18px 14px; }} .panel {{ padding:24px; }} .measurements {{ grid-template-columns:repeat(2,1fr); }} }}
    @media(max-width:460px) {{ form {{ display:grid; }} .measurements {{ grid-template-columns:1fr 1fr; }} }}
  </style>
</head>
<body>
  <div class="wrap">
    <aside class="sidebar">
      <div class="brand">Gooddegg Management</div>
      <div class="menu-title">Menu</div>
      <a class="menu-item" href="/"><span class="menu-icon">F</span><span>FOA Inventory</span></a>
      <a class="menu-item" href="/butler"><span class="menu-icon">B</span><span>Butler Inventory</span></a>
      <a class="menu-item active" href="/shipping"><span class="menu-icon">S</span><span>Shipping Info</span></a>
      <a class="menu-item" href="/order-notes"><span class="menu-icon">N</span><span>Order Notes</span></a>
    </aside>
    <main>
      <div class="panel">
        <span class="label">FOA Item Master</span>
        <h1>Shipping Info</h1>
        <p class="intro">Enter an exact SKU to view package length, width, height and weight. Group SKUs also show their component packages.</p>
        <form method="get" action="/shipping">
          <input type="search" name="sku" value="{html.escape(query)}" placeholder="e.g. FOA7730RD-BED" aria-label="SKU" autofocus required>
          <button type="submit">Find SKU</button>
        </form>
        {error_html}
        {result_html}
      </div>
    </main>
  </div>
</body>
</html>"""
    return html_body.encode("utf-8")


def render_order_notes_page(message: str = "", message_kind: str = "") -> bytes:
    notes: list[dict[str, str]] = []
    try:
        notes = load_order_notes()
    except Exception as exc:
        message = str(exc)
        message_kind = "error"

    note_cards = []
    for note in notes:
        note_id = html.escape(str(note.get("id", "")), quote=True)
        order_number = html.escape(str(note.get("order_number", "")), quote=True)
        note_text = html.escape(str(note.get("note", "")))
        updated_at = html.escape(str(note.get("updated_at", "")))
        note_cards.append(
            f"""<article class="note-card">
              <form method="post" action="/order-notes">
                <input type="hidden" name="action" value="save">
                <input type="hidden" name="note_id" value="{note_id}">
                <div class="note-head">
                  <input class="order-input" type="text" name="order_number" value="{order_number}" placeholder="Order number / title" maxlength="200">
                  <span class="updated">Updated {updated_at}</span>
                </div>
                <textarea name="note" placeholder="Write order note..." maxlength="10000" required>{note_text}</textarea>
                <div class="note-actions">
                  <button class="save-button" type="submit">Save Note</button>
                  <button class="delete-button" type="submit" name="action" value="delete" formnovalidate onclick="return confirm('Delete this note?')">Delete</button>
                </div>
              </form>
            </article>"""
        )

    notes_html = "".join(note_cards)
    if not notes_html:
        notes_html = '<div class="empty-state" id="emptyState"><strong>No order notes yet.</strong><span>Click + Add to create your first note.</span></div>'

    message_html = ""
    if message:
        message_html = f'<div class="message {html.escape(message_kind)}">{html.escape(message)}</div>'

    html_body = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Order Notes | Gooddegg Management</title>
  <style>
    :root {{ --panel:rgba(255,255,255,.86); --text:#17201b; --muted:#66736c; --line:rgba(23,32,27,.09); --accent:#163f31; --soft:#edf3ef; --danger:#a33d36; }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; font-family:"Segoe UI",Tahoma,sans-serif; color:var(--text); background:radial-gradient(circle at top left,rgba(22,63,49,.08),transparent 24%),linear-gradient(180deg,#f8faf8,#f3f5f2); }}
    .wrap {{ min-height:100vh; display:grid; grid-template-columns:230px minmax(0,850px); justify-content:center; gap:24px; padding:32px 18px; }}
    .sidebar {{ padding:24px 14px; border:1px solid rgba(255,255,255,.7); border-radius:24px; background:rgba(255,255,255,.7); box-shadow:0 18px 50px rgba(28,42,35,.06); }}
    .brand {{ padding:8px 12px 24px; color:var(--accent); font-size:1.05rem; font-weight:700; }}
    .menu-title {{ padding:0 12px 10px; color:var(--muted); font-size:.72rem; font-weight:700; letter-spacing:.1em; text-transform:uppercase; }}
    .menu-item {{ display:flex; align-items:center; gap:10px; padding:12px; border-radius:14px; color:var(--accent); font-weight:600; text-decoration:none; }}
    .menu-item.active {{ background:var(--soft); }}
    .menu-icon {{ display:grid; width:28px; height:28px; place-items:center; border-radius:9px; background:var(--accent); color:white; font-size:.78rem; }}
    main {{ min-width:0; }}
    .panel {{ background:var(--panel); border:1px solid rgba(255,255,255,.65); border-radius:28px; padding:34px; box-shadow:0 18px 50px rgba(28,42,35,.08); }}
    .page-head {{ display:flex; justify-content:space-between; align-items:flex-start; gap:18px; }}
    .label {{ display:inline-flex; padding:7px 12px; border-radius:999px; background:var(--soft); color:var(--accent); font-size:.76rem; font-weight:700; letter-spacing:.06em; text-transform:uppercase; }}
    h1 {{ margin:16px 0 8px; font-size:clamp(2rem,5vw,3rem); line-height:1.02; letter-spacing:-.04em; }}
    .intro {{ margin:0; color:var(--muted); line-height:1.6; }}
    button {{ border:0; cursor:pointer; font:inherit; font-weight:700; }}
    .add-button {{ flex:none; padding:13px 20px; border-radius:15px; background:var(--accent); color:white; box-shadow:0 10px 24px rgba(22,63,49,.16); }}
    .notes {{ display:grid; gap:14px; margin-top:24px; }}
    .note-card,.new-note,.empty-state {{ padding:20px; border:1px solid var(--line); border-radius:20px; background:rgba(255,255,255,.78); }}
    .new-note {{ display:none; border-color:rgba(22,63,49,.25); box-shadow:0 12px 28px rgba(22,63,49,.08); }}
    .new-note.open {{ display:block; }}
    .note-head {{ display:flex; align-items:center; justify-content:space-between; gap:14px; margin-bottom:12px; }}
    .order-input {{ flex:1; min-width:0; border:0; border-bottom:1px solid var(--line); padding:9px 2px; background:transparent; color:var(--text); font-size:1.05rem; font-weight:700; outline:none; }}
    .order-input:focus {{ border-color:var(--accent); }}
    .updated {{ flex:none; color:var(--muted); font-size:.76rem; }}
    textarea {{ display:block; width:100%; min-height:130px; resize:vertical; border:1px solid var(--line); border-radius:15px; padding:14px; background:#fff; color:var(--text); font:inherit; line-height:1.55; outline:none; }}
    textarea:focus {{ border-color:rgba(22,63,49,.45); box-shadow:0 0 0 4px rgba(22,63,49,.07); }}
    .note-actions {{ display:flex; gap:9px; justify-content:flex-end; margin-top:12px; }}
    .save-button,.cancel-button,.delete-button {{ padding:10px 15px; border-radius:12px; }}
    .save-button {{ background:var(--accent); color:white; }}
    .cancel-button {{ background:var(--soft); color:var(--accent); }}
    .delete-button {{ margin-right:auto; background:#fff2f1; color:var(--danger); }}
    .empty-state {{ display:grid; gap:5px; place-items:center; min-height:145px; color:var(--muted); text-align:center; }}
    .message {{ margin-top:20px; padding:13px 16px; border-radius:14px; background:var(--soft); color:var(--accent); }}
    .message.error {{ background:#fff3f2; color:#9a3b35; }}
    @media(max-width:760px) {{ .wrap {{ grid-template-columns:1fr; padding:18px 14px; }} .panel {{ padding:24px; }} .sidebar {{ min-height:auto; }} }}
    @media(max-width:520px) {{ .page-head,.note-head {{ display:grid; }} .add-button {{ width:100%; }} .updated {{ order:-1; }} }}
  </style>
</head>
<body>
  <div class="wrap">
    <aside class="sidebar">
      <div class="brand">Gooddegg Management</div>
      <div class="menu-title">Menu</div>
      <a class="menu-item" href="/"><span class="menu-icon">F</span><span>FOA Inventory</span></a>
      <a class="menu-item" href="/butler"><span class="menu-icon">B</span><span>Butler Inventory</span></a>
      <a class="menu-item" href="/shipping"><span class="menu-icon">S</span><span>Shipping Info</span></a>
      <a class="menu-item active" href="/order-notes"><span class="menu-icon">N</span><span>Order Notes</span></a>
    </aside>
    <main>
      <div class="panel">
        <div class="page-head">
          <div><span class="label">Order Workspace</span><h1>Order Notes</h1><p class="intro">Keep order details, reminders and follow-up notes in one place.</p></div>
          <button class="add-button" type="button" id="addNoteButton">+ Add</button>
        </div>
        {message_html}
        <section class="notes">
          <article class="new-note" id="newNote">
            <form method="post" action="/order-notes">
              <input type="hidden" name="action" value="save">
              <input class="order-input" type="text" name="order_number" placeholder="Order number / title" maxlength="200" id="newOrderNumber">
              <textarea name="note" placeholder="Write order note..." maxlength="10000" required></textarea>
              <div class="note-actions"><button class="save-button" type="submit">Save Note</button><button class="cancel-button" type="button" id="cancelNoteButton">Cancel</button></div>
            </form>
          </article>
          {notes_html}
        </section>
      </div>
    </main>
  </div>
  <script>
    const addButton = document.getElementById('addNoteButton');
    const cancelButton = document.getElementById('cancelNoteButton');
    const newNote = document.getElementById('newNote');
    const orderInput = document.getElementById('newOrderNumber');
    addButton.addEventListener('click', () => {{ newNote.classList.add('open'); orderInput.focus(); }});
    cancelButton.addEventListener('click', () => {{ newNote.classList.remove('open'); newNote.querySelector('form').reset(); }});
  </script>
</body>
</html>"""
    return html_body.encode("utf-8")


def render_orders_page(message: str = "", message_kind: str = "", orders: list[dict[str, str]] | None = None) -> bytes:
    settings = load_shopify_settings()
    token_ready = bool(settings["admin_access_token"])
    escaped_message = html.escape(message)
    badge = "Connected" if token_ready else "Not connected"
    badge_class = "ok" if token_ready else "warn"
    rows = ""
    for order in orders or []:
        rows += (
            "<tr>"
            f"<td>{html.escape(order['name'])}</td>"
            f"<td>{html.escape(order['customer'])}</td>"
            f"<td>{html.escape(order['total_price'])} {html.escape(order['currency'])}</td>"
            f"<td>{html.escape(order['financial_status'])}</td>"
            f"<td>{html.escape(order['fulfillment_status'])}</td>"
            f"<td>{html.escape(order['created_at'])}</td>"
            "</tr>"
        )

    if not rows:
        rows = '<tr><td colspan="6">No orders loaded yet.</td></tr>'

    html_body = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Shopify Orders</title>
  <style>
    :root {{
      --bg: #f6f4ef;
      --panel: rgba(255,255,255,0.88);
      --text: #221d16;
      --muted: #6d655d;
      --line: rgba(34,29,22,0.1);
      --accent: #6c4b2b;
      --accent-soft: #f3ebe1;
      --ok: #2b6b3f;
      --warn: #9a5a18;
      --error: #9b3b34;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: "Segoe UI", Tahoma, sans-serif;
      color: var(--text);
      background:
        radial-gradient(circle at top right, rgba(108, 75, 43, 0.10), transparent 25%),
        linear-gradient(180deg, #fbfaf7 0%, #f3efe8 100%);
    }}
    .wrap {{
      min-height: 100vh;
      padding: 28px 18px 40px;
    }}
    .panel {{
      width: min(1080px, 100%);
      margin: 0 auto;
      background: var(--panel);
      border: 1px solid rgba(255,255,255,0.7);
      border-radius: 28px;
      padding: 28px;
      box-shadow: 0 18px 50px rgba(40, 30, 18, 0.08);
      backdrop-filter: blur(10px);
    }}
    .topbar {{
      display: flex;
      justify-content: space-between;
      gap: 12px;
      align-items: center;
      flex-wrap: wrap;
      margin-bottom: 18px;
    }}
    .topbar a {{
      color: var(--accent);
      text-decoration: none;
      font-weight: 600;
    }}
    .badge {{
      display: inline-flex;
      align-items: center;
      padding: 8px 12px;
      border-radius: 999px;
      background: var(--accent-soft);
      font-size: 0.85rem;
      font-weight: 700;
    }}
    .badge.ok {{ color: var(--ok); }}
    .badge.warn {{ color: var(--warn); }}
    h1 {{
      margin: 0 0 10px;
      font-size: clamp(2rem, 5vw, 3rem);
      letter-spacing: -0.04em;
    }}
    .intro {{
      margin: 0 0 18px;
      color: var(--muted);
      line-height: 1.6;
    }}
    .grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(240px, 1fr));
      gap: 14px;
      margin-bottom: 18px;
    }}
    .card {{
      border: 1px solid var(--line);
      border-radius: 18px;
      padding: 16px;
      background: rgba(255,255,255,0.75);
    }}
    .card strong {{
      display: block;
      margin-bottom: 6px;
      font-size: 0.9rem;
    }}
    .card span {{
      color: var(--muted);
      word-break: break-word;
    }}
    .actions {{
      display: flex;
      gap: 12px;
      flex-wrap: wrap;
      margin-bottom: 18px;
    }}
    .button {{
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-height: 48px;
      padding: 0 18px;
      border-radius: 16px;
      background: var(--accent);
      color: white;
      text-decoration: none;
      font-weight: 700;
      box-shadow: 0 10px 24px rgba(108, 75, 43, 0.18);
    }}
    .button.secondary {{
      background: white;
      color: var(--accent);
      border: 1px solid var(--line);
      box-shadow: none;
    }}
    .note {{
      margin-bottom: 18px;
      padding: 14px 16px;
      border-radius: 16px;
      border: 1px solid var(--line);
      background: rgba(255,255,255,0.7);
      color: var(--muted);
    }}
    .note.error {{ color: var(--error); }}
    .note.ok {{ color: var(--ok); }}
    table {{
      width: 100%;
      border-collapse: collapse;
      overflow: hidden;
      border-radius: 18px;
      background: white;
    }}
    th, td {{
      padding: 14px 12px;
      text-align: left;
      border-bottom: 1px solid var(--line);
      font-size: 0.95rem;
      vertical-align: top;
    }}
    th {{
      background: #f8f3ec;
      font-size: 0.82rem;
      text-transform: uppercase;
      letter-spacing: 0.06em;
      color: var(--muted);
    }}
    @media (max-width: 720px) {{
      .panel {{ padding: 22px; border-radius: 22px; }}
      table, thead, tbody, th, td, tr {{ display: block; }}
      thead {{ display: none; }}
      tr {{
        border-bottom: 1px solid var(--line);
        padding: 8px 0;
      }}
      td {{
        border: 0;
        padding: 8px 0;
      }}
    }}
  </style>
</head>
<body>
  <div class="wrap">
    <div class="panel">
      <div class="topbar">
        <a href="/">Back to inventory app</a>
        <span class="badge {badge_class}">{badge}</span>
      </div>
      <h1>Shopify order list</h1>
      <p class="intro">This page uses your existing Gooddegg Shopify app credentials. Connect once, then load the latest {SHOPIFY_ORDERS_LIMIT} orders directly from Shopify Admin API.</p>
      <div class="grid">
        <div class="card"><strong>Shop Domain</strong><span>{html.escape(settings["shop_domain"] or "Not found")}</span></div>
        <div class="card"><strong>API Version</strong><span>{html.escape(settings["api_version"])}</span></div>
        <div class="card"><strong>Client ID</strong><span>{html.escape(settings["client_id"] or "Not found")}</span></div>
        <div class="card"><strong>Token Status</strong><span>{"Saved in local .env" if token_ready else "Not saved yet"}</span></div>
      </div>
      <div class="actions">
        <a class="button" href="/shopify/connect">Connect Shopify</a>
        <a class="button secondary" href="/orders?refresh=1">Load latest orders</a>
      </div>
      <div class="note {html.escape(message_kind)}">{escaped_message or "If token is missing, click Connect Shopify first. After successful connection, the access token will be saved in this project's .env file."}</div>
      <table>
        <thead>
          <tr>
            <th>Order</th>
            <th>Customer</th>
            <th>Total</th>
            <th>Payment</th>
            <th>Fulfillment</th>
            <th>Created</th>
          </tr>
        </thead>
        <tbody>{rows}</tbody>
      </table>
    </div>
  </div>
</body>
</html>"""
    return html_body.encode("utf-8")


def parse_uploaded_files(
    handler: BaseHTTPRequestHandler,
) -> tuple[list[tuple[str, bytes]], tuple[str, bytes] | None]:
    content_type = handler.headers.get("Content-Type", "")
    if "multipart/form-data" not in content_type:
        raise ValueError("Expected a multipart form upload.")

    content_length = int(handler.headers.get("Content-Length", "0"))
    if content_length <= 0:
        raise ValueError("No uploaded file data was received.")

    body = handler.rfile.read(content_length)
    message = BytesParser(policy=default).parsebytes(
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8") + body
    )

    files: list[tuple[str, bytes]] = []
    asset_file: tuple[str, bytes] | None = None
    for part in message.iter_parts():
        if part.get_content_disposition() != "form-data":
            continue
        field_name = part.get_param("name", header="content-disposition")
        filename = part.get_filename()
        if not filename:
            continue

        content = part.get_payload(decode=True) or b""
        if not content:
            continue

        if field_name == "inventory_files":
            files.append((filename, content))
        elif field_name == "butler_inventory_file":
            files.append((filename, content))
        elif field_name == "asset_file":
            asset_file = (filename, content)

    return files, asset_file


class InventoryHandler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:
        parsed_url = parse.urlparse(self.path)

        if parsed_url.path in ("/", "/butler"):
            body = render_home(mode="butler" if parsed_url.path == "/butler" else "foa")
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        if parsed_url.path == "/shipping":
            query = parse.parse_qs(parsed_url.query).get("sku", [""])[-1].strip()
            body = render_shipping_page(query=query)
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        if parsed_url.path == "/order-notes":
            query = parse.parse_qs(parsed_url.query)
            message = ""
            if "saved" in query:
                message = "Order note saved."
            elif "deleted" in query:
                message = "Order note deleted."
            body = render_order_notes_page(message=message, message_kind="ok" if message else "")
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        if parsed_url.path == "/orders":
            query = parse.parse_qs(parsed_url.query)
            message = ""
            message_kind = ""
            orders: list[dict[str, str]] = []
            if "connected" in query:
                message = "Shopify connected successfully. You can load the latest orders now."
                message_kind = "ok"

            if "refresh" in query:
                try:
                    orders = fetch_shopify_orders(load_shopify_settings())
                    message = f"Loaded {len(orders)} orders from Shopify."
                    message_kind = "ok"
                except Exception as exc:
                    message = str(exc)
                    message_kind = "error"

            body = render_orders_page(message=message, message_kind=message_kind, orders=orders)
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        if parsed_url.path == "/shopify/connect":
            settings = load_shopify_settings()
            if not settings["shop_domain"] or not settings["client_id"] or not settings["client_secret"]:
                body = render_orders_page(
                    message="Shopify domain, client ID, or client secret was not found in config.",
                    message_kind="error",
                )
                self.send_response(HTTPStatus.BAD_REQUEST)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return

            state = hashlib.sha256(os.urandom(32)).hexdigest()
            SHOPIFY_RUNTIME_STATE[SHOPIFY_STATE_KEY] = state
            self.send_response(HTTPStatus.FOUND)
            self.send_header("Location", build_shopify_authorization_url(settings, state))
            self.end_headers()
            return

        if parsed_url.path == "/shopify/callback":
            query = {key: values[-1] for key, values in parse.parse_qs(parsed_url.query).items()}
            settings = load_shopify_settings()

            try:
                expected_state = SHOPIFY_RUNTIME_STATE.get(SHOPIFY_STATE_KEY, "")
                if not expected_state or query.get("state", "") != expected_state:
                    raise ValueError("OAuth state mismatch. Please try connecting again.")
                if not validate_shopify_hmac(query, settings["client_secret"]):
                    raise ValueError("Invalid Shopify callback signature.")
                code = query.get("code", "").strip()
                if not code:
                    raise ValueError("Shopify callback did not include an authorization code.")

                token = exchange_shopify_code_for_token(settings, code)
                save_env_value(ENV_PATH, SHOPIFY_TOKEN_KEY, token)
                SHOPIFY_RUNTIME_STATE.pop(SHOPIFY_STATE_KEY, None)
                self.send_response(HTTPStatus.FOUND)
                self.send_header("Location", "/orders?connected=1")
                self.end_headers()
                return
            except Exception as exc:
                body = render_orders_page(message=str(exc), message_kind="error")
                self.send_response(HTTPStatus.BAD_REQUEST)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return

        self.send_error(HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:
        parsed_path = parse.urlparse(self.path).path
        if parsed_path == "/order-notes":
            try:
                content_type = self.headers.get("Content-Type", "")
                if "application/x-www-form-urlencoded" not in content_type:
                    raise ValueError("Invalid order note submission.")
                content_length = int(self.headers.get("Content-Length", "0"))
                if content_length <= 0 or content_length > 1_000_000:
                    raise ValueError("Order note submission is empty or too large.")
                body = self.rfile.read(content_length).decode("utf-8")
                fields = parse.parse_qs(body, keep_blank_values=True)
                action = fields.get("action", ["save"])[-1]
                note_id = fields.get("note_id", [""])[-1].strip()

                if action == "delete":
                    if not note_id:
                        raise ValueError("Order note was not found.")
                    delete_order_note(note_id)
                    location = "/order-notes?deleted=1"
                else:
                    save_order_note(
                        note_id,
                        fields.get("order_number", [""])[-1],
                        fields.get("note", [""])[-1],
                    )
                    location = "/order-notes?saved=1"

                self.send_response(HTTPStatus.SEE_OTHER)
                self.send_header("Location", location)
                self.end_headers()
                return
            except Exception as exc:
                body = render_order_notes_page(message=str(exc), message_kind="error")
                self.send_response(HTTPStatus.BAD_REQUEST)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return

        try:
            files, asset_file = parse_uploaded_files(self)
            is_butler = parse.urlparse(self.path).path == "/butler"

            if is_butler and len(files) != 1:
                raise ValueError("Please upload one Butler inventory CSV file.")
            if not is_butler and len(files) < 2:
                raise ValueError("Please upload at least two inventory files.")

            totals = merge_inventory(files)
            if not totals:
                raise ValueError("No SKU records were found in the uploaded files.")

            if is_butler:
                template_path = next(
                    (
                        path
                        for path in (BUTLER_EXTERNAL_TEMPLATE, BUTLER_ASSET_TEMPLATE)
                        if os.path.exists(path)
                    ),
                    None,
                )
                if template_path is None:
                    raise ValueError("Butler base file merged_inventory_export.csv was not found.")
                output_bytes = build_updated_inventory_csv(
                    totals,
                    template_path=template_path,
                    target_location=None,
                    minimum_quantity=MIN_EXPORT_QTY,
                    normalize_numeric_skus=True,
                )
                today = datetime.now()
                output_name = f"Butler_inventory_{today.day}_{today.strftime('%b').lower()}_{today.year}.csv"
            else:
                output_bytes = build_updated_inventory_csv(totals, asset_file=asset_file)
                today = datetime.now()
                output_name = (
                    f"Furniture_of_America_inventory_{today.day}_"
                    f"{today.strftime('%b').lower()}_{today.year}.csv"
                )

            self.send_response(HTTPStatus.OK)
            self.send_header(
                "Content-Type",
                "text/csv; charset=utf-8",
            )
            self.send_header("Content-Disposition", f'attachment; filename="{output_name}"')
            self.send_header("Content-Length", str(len(output_bytes)))
            self.end_headers()
            self.wfile.write(output_bytes)
        except Exception as exc:
            body = render_home(str(exc), mode="butler" if self.path == "/butler" else "foa")
            self.send_response(HTTPStatus.BAD_REQUEST)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


def run_server() -> None:
    server = ThreadingHTTPServer((HOST, PORT), InventoryHandler)
    print(f"Inventory app running at http://{HOST}:{PORT}")
    threading.Timer(0.8, lambda: webbrowser.open(f"http://{HOST}:{PORT}")).start()
    server.serve_forever()


if __name__ == "__main__":
    run_server()
